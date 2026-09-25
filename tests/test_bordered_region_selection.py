"""Which bordered rectangle in a spreadsheet is the table of movements.

WHY IT EXISTS
    The question that started this was whether images in an Excel file break
    the reading of the transaction list. They do not: openpyxl reads cell
    values and an image is a drawing, so a file with a logo is read exactly
    like the same file without one. The first case below records that, so the
    answer stays answered.

    Looking for it found something else. The detector took the topmost
    bordered rectangle and stopped there. A letterhead drawn with cell borders
    - logo, period, account holder, the shape of a great many real statements
    - was therefore read as the table, and the header row came out as the name
    of the bank. It won even when the movements below were bordered too,
    because the scan never reached them: the one layout the border detection
    exists to serve was the one it lost.

    Worse, it said it was certain, and downstream that certainty disables the
    second look at the header rows. A wrong answer that admits doubt is
    recoverable; a wrong answer that does not is what ships.

WHAT DECIDES IT NOW
    A bordered rectangle whose cells are mostly empty is a frame around
    letterhead, not a table. Measured on these fixtures: the framed logo block
    fills 0.33 of its cells, a real bordered movements table fills 1.00.
"""

from __future__ import annotations

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.styles import Border, Side  # noqa: E402

from core.normalizer import detect_bordered_region, detect_skip_rows  # noqa: E402

HEADER_ROWS = 4  # three lines of letterhead, one blank, then the column names

MOVEMENTS = [
    ("01/01/2026", "POS SUPERMERCATO", -45.20),
    ("03/01/2026", "BONIFICO STIPENDIO", 2100.00),
    ("05/01/2026", "ADDEBITO UTENZE", -88.10),
    ("09/01/2026", "POS RISTORANTE", -32.00),
    ("12/01/2026", "PRELIEVO ATM", -100.00),
    ("15/01/2026", "POS FARMACIA", -18.50),
]


def _statement(path, *, image=False, frame_logo=False, frame_table=False) -> bytes:
    """A statement with letterhead, then the movements. One variable at a time."""
    thin = Border(*[Side(style="thin")] * 4)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimenti"

    ws["A1"] = "Banca di Prova S.p.A."
    ws["A2"] = "Estratto conto"
    ws["A3"] = "Periodo: 01/01/2026 - 31/01/2026"
    if frame_logo:
        for r in range(1, 4):
            for c in range(1, 4):
                ws.cell(row=r, column=c).border = thin

    for c, name in enumerate(["Data", "Descrizione", "Importo"], start=1):
        ws.cell(row=5, column=c, value=name)
    for i, (date, description, amount) in enumerate(MOVEMENTS, start=6):
        ws.cell(row=i, column=1, value=date)
        ws.cell(row=i, column=2, value=description)
        ws.cell(row=i, column=3, value=amount)
    if frame_table:
        for r in range(5, 6 + len(MOVEMENTS)):
            for c in range(1, 4):
                ws.cell(row=r, column=c).border = thin

    if image:
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage

        logo = path.parent / "logo.png"
        PILImage.new("RGB", (160, 80), (200, 30, 30)).save(logo)
        ws.add_image(XLImage(str(logo)), "E1")

    wb.save(path)
    return path.read_bytes()


LAYOUTS = {
    "plain": {},
    "with an image": {"image": True},
    "with a framed letterhead": {"frame_logo": True},
    "with a framed letterhead and an image": {"frame_logo": True, "image": True},
    "with a bordered table": {"frame_table": True},
    "with both bordered": {"frame_logo": True, "frame_table": True},
}


@pytest.mark.parametrize("layout", list(LAYOUTS), ids=list(LAYOUTS))
def test_the_movements_start_where_they_start(tmp_path, layout):
    """Every layout, one expected answer: the header is row 4.

    Parameterised rather than written out, because the point is that the
    answer does not depend on the decoration.
    """
    raw = _statement(tmp_path / "statement.xlsx", **LAYOUTS[layout])

    skip_rows, _certain, _region = detect_skip_rows(raw, "statement.xlsx")

    assert skip_rows == HEADER_ROWS, (
        f"a statement {layout} was read as starting at row {skip_rows}"
    )


def test_an_image_changes_nothing_at_all(tmp_path):
    """The original question, kept as its own answer.

    A drawing is not a cell, so the two files are read identically. Recorded
    so nobody has to work it out again from first principles.
    """
    plain = _statement(tmp_path / "plain.xlsx")
    with_logo = _statement(tmp_path / "logo.xlsx", image=True)

    assert detect_skip_rows(plain, "plain.xlsx") == detect_skip_rows(with_logo, "logo.xlsx")


def test_a_frame_around_letterhead_is_not_offered_as_a_table(tmp_path):
    """Saying nothing beats being confidently wrong.

    Returning None here sends the caller to the density scan, which reads this
    layout correctly. Returning the frame instead came with certainty, and
    certainty is what switches off the second look.
    """
    raw = _statement(tmp_path / "statement.xlsx", frame_logo=True)

    assert detect_bordered_region(raw, "statement.xlsx") is None


def test_a_real_bordered_table_is_still_found(tmp_path):
    """The layout the border detection exists to serve."""
    raw = _statement(tmp_path / "statement.xlsx", frame_table=True)

    region = detect_bordered_region(raw, "statement.xlsx")

    assert region is not None, "a bordered movements table was not recognised"
    first_row, last_row, _first_col, _last_col = region
    assert first_row == HEADER_ROWS, "the region does not start at the column names"
    assert last_row == HEADER_ROWS + len(MOVEMENTS)


def test_the_table_wins_over_the_letterhead_when_both_are_bordered(tmp_path):
    """The case that made the old rule lose on its own home ground."""
    raw = _statement(tmp_path / "statement.xlsx", frame_logo=True, frame_table=True)

    region = detect_bordered_region(raw, "statement.xlsx")

    assert region is not None
    assert region[0] == HEADER_ROWS, (
        f"the letterhead was chosen over the movements: region starts at {region[0]}"
    )
